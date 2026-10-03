"""
附件解析器
解析 zip 压缩包中的 doc/docx/xlsx/pdf 等附件文件，提取文本和表格
"""
import io
import re
import threading
import zipfile
from pathlib import Path
from typing import Dict, List, Tuple

from models.schemas import AnnouncementContent


_OCR_ENGINE = None
_OCR_INIT_LOCK = threading.Lock()
_OCR_RUN_LOCK = threading.Lock()


def _get_ocr_engine():
    """惰性初始化 OCR 引擎。

    主流程会并发解析附件，所以全进程只初始化一份模型，
    避免每篇公告重复加载模型并瞬间占满内存。
    """
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        with _OCR_INIT_LOCK:
            if _OCR_ENGINE is None:
                from rapidocr import RapidOCR
                _OCR_ENGINE = RapidOCR(params={"Global.log_level": "warning"})
    return _OCR_ENGINE


# ============================================================
# PDF 文本层「字符重复」修正
# ============================================================
# 现象：部分 PDF 生成器会把每个字形写入两层，导致提取出的文本
#       每个字都重复一遍，例如「绥化市公安局」→「绥绥化化市市公公安安局局」。
#       这类文本白白消耗一倍 token，也可能干扰大模型提取。
#
# 为什么按「词块」而不是按「整篇」判定：
#       实测中存在混合情况——同一份 PDF 里标题段是重复的、正文是正常的
#       （见 t20260202_26139917 的报价明细附件）。若按整篇比例判定，
#       这类文件的重复比例被正文稀释到阈值以下，会整篇跳过、等于没修。
#       改为逐个空白分隔的词块独立判定，才能精确命中重复段。
#
# 为什么不能无脑用 (.)\1 -> \1：
#       那会毁掉「谢谢」「常常」等正常叠词。这里要求词块足够长
#       且成对位置高度一致才折叠，正常文本几乎不可能触发。
# ============================================================

# 词块被判定为「重复」所需的最少字符数（排除空白后）
_DOUBLED_TOKEN_MIN_LEN = 12
# 词块中交替成对位置字符相同的最低占比
_DOUBLED_TOKEN_RATIO = 0.9
# 词块中最少需要的不同字符数
# 用于挡掉「谢谢谢谢常常常常…」「AAAA…」这类「只由少数几种字反复叠加」的
# 正常文本——它们同样满足成对相同，但并非 PDF 双层字形造成的重复。
_DOUBLED_MIN_DISTINCT = 4


def _doubled_pair_ratio(chars: list, offset: int = 0) -> float:
    """计算从 offset 起「交替成对位置字符相同」的比例"""
    sub = chars[offset:]
    pairs = len(sub) // 2
    if pairs == 0:
        return 0.0
    hit = sum(1 for i in range(0, pairs * 2, 2) if sub[i] == sub[i + 1])
    return hit / pairs


def is_doubled_token(token: str) -> bool:
    """判断单个词块是否整体呈「每字符重复一次」的特征"""
    chars = [c for c in token if not c.isspace()]
    if len(chars) < _DOUBLED_TOKEN_MIN_LEN:
        return False
    if len(set(chars)) < _DOUBLED_MIN_DISTINCT:
        return False
    # 起始位置可能整体错位一位，两种对齐取较大值，避免漏判
    best = max(_doubled_pair_ratio(chars, 0), _doubled_pair_ratio(chars, 1))
    return best >= _DOUBLED_TOKEN_RATIO


def _collapse_pairs(token: str) -> str:
    """把词块中相邻且相同的非空白字符折叠为一个（空白原样保留）"""
    out = []
    i = 0
    n = len(token)
    while i < n:
        ch = token[i]
        if ch.isspace():
            out.append(ch)
            i += 1
            continue
        if i + 1 < n and token[i + 1] == ch:
            i += 2
        else:
            i += 1
        out.append(ch)
    return "".join(out)


def undouble_text(text: str) -> str:
    """
    逐词块修正重复字符，对正常文本不做任何改动

    可安全地无条件调用：只处理被 is_doubled_token() 判定为重复的词块。
    """
    parts = re.split(r"(\s+)", text)
    for idx, part in enumerate(parts):
        if part and not part.isspace() and is_doubled_token(part):
            parts[idx] = _collapse_pairs(part)
    return "".join(parts)


class AttachmentParser:
    """解析公告附件（zip 包中的多格式文件）"""

    SUPPORTED_EXTENSIONS = {".doc", ".docx", ".xlsx", ".xls", ".pdf", ".pptx", ".txt"}

    def __init__(self):
        # 记录本次被修正过「字符重复」的文件名，供数据质量分析统计
        self.dedoubled_files: List[str] = []

    def parse_zip(self, zip_path: str, announcement_id: str = "") -> Tuple[Dict[str, str], Dict[str, List]]:
        """
        解析 zip 附件包
        返回: (附件文本字典, 附件表格字典)
        """
        texts = {}
        tables = {}

        path = Path(zip_path)
        if not path.exists():
            return texts, tables

        try:
            with zipfile.ZipFile(zip_path, "r") as zf:
                for file_info in zf.infolist():
                    if file_info.is_dir():
                        continue

                    ext = Path(file_info.filename).suffix.lower()
                    if ext not in self.SUPPORTED_EXTENSIONS:
                        continue

                    try:
                        file_data = zf.read(file_info.filename)
                        text, table = self._parse_file(file_data, ext, file_info.filename)
                        if text:
                            texts[file_info.filename] = text
                        if table:
                            tables[file_info.filename] = table
                    except Exception:
                        # 静默跳过无法解析的附件（加密/格式不符等），HTML 中通常已有主要信息
                        continue
        except zipfile.BadZipFile:
            pass

        return texts, tables

    def parse_files(self, file_paths: List[str]) -> Tuple[Dict[str, str], Dict[str, List]]:
        """直接解析多个文件（非 zip）"""
        texts = {}
        tables = {}

        for fp in file_paths:
            path = Path(fp)
            if not path.exists():
                continue
            ext = path.suffix.lower()
            if ext not in self.SUPPORTED_EXTENSIONS:
                continue

            try:
                file_data = path.read_bytes()
                text, table = self._parse_file(file_data, ext, path.name)
                if text:
                    texts[path.name] = text
                if table:
                    tables[path.name] = table
            except Exception as e:
                print(f"  [警告] 解析文件 {fp} 失败: {e}")

        return texts, tables

    def _parse_file(self, file_data: bytes, ext: str, filename: str) -> Tuple[str, List]:
        """根据扩展名解析单个文件，返回 (文本, 表格)"""
        if ext == ".docx":
            text, tables = self._parse_docx(file_data)
        elif ext == ".doc":
            text, tables = self._parse_doc(file_data)
        elif ext == ".xlsx":
            text, tables = self._parse_excel(file_data)
        elif ext == ".xls":
            text, tables = self._parse_xls(file_data)
        elif ext == ".pdf":
            text, tables = self._parse_pdf(file_data, filename)
        elif ext == ".pptx":
            text, tables = self._parse_pptx(file_data)
        elif ext == ".txt":
            text, tables = self._parse_txt(file_data)
        else:
            text, tables = "", []

        # 修正 PDF 文本层重复字符（逐词块判定，正常文本不受影响）
        if text and ext == ".pdf":
            cleaned = undouble_text(text)
            if cleaned != text:
                text = cleaned
                self.dedoubled_files.append(filename)

        return text, tables

    def _parse_docx(self, file_data: bytes) -> Tuple[str, List]:
        """解析 Word .docx 文件"""
        from docx import Document

        doc = Document(io.BytesIO(file_data))

        # 提取段落文本
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())

        # 提取表格
        tables = []
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            if table_data:
                tables.append(table_data)
                # 表格内容也加入文本
                for row in table_data:
                    text_parts.append(" | ".join(row))

        return "\n".join(text_parts), tables

    def _parse_doc(self, file_data: bytes) -> Tuple[str, List]:
        """解析后缀为 .doc 的文件，先识别真实容器类型。"""
        # 数据集中有一部分文件只是被错误命名为 .doc，
        # 实际上是 DOCX（ZIP 容器）或 HTML。
        if file_data.startswith(b"PK\x03\x04"):
            return self._parse_docx(file_data)

        prefix = file_data[:4096].lstrip().lower()
        if prefix.startswith(b"<"):
            from bs4 import BeautifulSoup
            html, _ = self._parse_txt(file_data)
            soup = BeautifulSoup(html, "lxml")
            return soup.get_text("\n", strip=True), []

        # 真正的 Word 97–2003 文件是 OLE Compound File。
        if file_data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            try:
                from legacy_doc import extract_text
                result = extract_text(file_data)
                return (result.text or "").strip(), []
            except Exception:
                # 某些 WPS/非标准 OLE 文件的扇区链有损坏，legacy-doc
                # 可能拒绝整个文件。这些文件中通常仍保留了连续的
                # UTF-16LE 正文，只抢救长可读片段，不做无条件二进制解码。
                return self._salvage_doc_unicode(file_data), []

        # 未知二进制格式不再强行解码，避免把乱码误报为有效正文。
        return "", []

    @staticmethod
    def _salvage_doc_unicode(file_data: bytes) -> str:
        """从损坏/非标准 OLE Word 中抢救连续 UTF-16LE 文本片段。"""
        decoded = file_data.decode("utf-16le", errors="ignore")
        readable_runs = re.findall(
            r"[\u4e00-\u9fffA-Za-z0-9，。；：！？、（）【】《》()\-_/ .]{8,}",
            decoded,
        )
        metadata = {
            "root entry",
            "summaryinformation",
            "documentsummaryinformation",
            "worddocument",
            "normal.dot",
            "objectpool",
        }
        cleaned = []
        for run in readable_runs:
            run = re.sub(r"\s+", " ", run).strip(" .")
            if not run or run.lower() in metadata:
                continue
            lower_run = run.lower()
            if (
                any(marker in lower_run for marker in ("normal.dot", "wps office", "microsoft office"))
                and len(re.findall(r"[\u4e00-\u9fff]", run)) < 20
            ):
                continue
            cleaned.append(run)
        text = "\n".join(cleaned)
        # 只有 OLE 元数据、无足够中文正文时仍视为解析失败。
        if len(re.findall(r"[\u4e00-\u9fff]", text)) < 20:
            return ""
        return text

    def _parse_excel(self, file_data: bytes) -> Tuple[str, List]:
        """解析 OOXML Excel（.xlsx）文件"""
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(file_data), read_only=True, data_only=True)

        text_parts = []
        all_tables = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_data = []

            for row in ws.iter_rows(values_only=True):
                row_data = [str(cell) if cell is not None else "" for cell in row]
                # 跳过全空行
                if any(cell.strip() for cell in row_data):
                    sheet_data.append(row_data)
                    text_parts.append(" | ".join(row_data))

            if sheet_data:
                all_tables.append(sheet_data)

        wb.close()
        return "\n".join(text_parts), all_tables

    def _parse_xls(self, file_data: bytes) -> Tuple[str, List]:
        """解析旧版二进制 Excel（.xls）文件。"""
        import xlrd

        book = xlrd.open_workbook(file_contents=file_data, on_demand=True)
        text_parts = []
        all_tables = []

        try:
            for sheet in book.sheets():
                sheet_data = []
                for row_idx in range(sheet.nrows):
                    row_data = []
                    for value in sheet.row_values(row_idx):
                        if isinstance(value, float) and value.is_integer():
                            value = int(value)
                        row_data.append(str(value).strip() if value not in (None, "") else "")
                    if any(row_data):
                        sheet_data.append(row_data)
                        text_parts.append(" | ".join(row_data))
                if sheet_data:
                    all_tables.append(sheet_data)
        finally:
            book.release_resources()

        return "\n".join(text_parts), all_tables

    @staticmethod
    def _select_pdf_page_indices(total_pages: int, filename: str = "") -> List[int]:
        """选取 PDF 页：短文档全读，长文档覆盖页首、中部与页尾。"""
        import config

        relevant = bool(re.search(config.PDF_RELEVANT_NAME_PATTERN, filename, re.IGNORECASE))
        max_pages = (
            config.PDF_RELEVANT_MAX_TEXT_PAGES if relevant else config.PDF_MAX_TEXT_PAGES
        )
        max_pages = max(1, max_pages)
        if total_pages <= max_pages:
            return list(range(total_pages))

        # 前半配额连续取首页；后半配额在剩余页面均匀取样，
        # 其中必然包含最后一页。
        head_count = max_pages // 2
        selected = list(range(head_count))
        sample_count = max_pages - head_count
        start = head_count
        span = total_pages - 1 - start
        if sample_count == 1:
            selected.append(total_pages - 1)
        else:
            selected.extend(
                round(start + span * i / (sample_count - 1))
                for i in range(sample_count)
            )
        return sorted(set(selected))

    @staticmethod
    def _ocr_pdf_page(pdf_document, page_index: int) -> str:
        """渲染单页并 OCR；任何失败都返回空串，不中断整篇公告。"""
        import config
        import numpy as np

        try:
            page = pdf_document[page_index]
            bitmap = page.render(scale=config.OCR_RENDER_SCALE)
            image = np.asarray(bitmap.to_pil().convert("RGB"))
            # RapidOCR/ONNX 不保证同一实例的并发推理安全，这里仅串行化 OCR 调用。
            with _OCR_RUN_LOCK:
                result = _get_ocr_engine()(image)
            texts = getattr(result, "txts", None) or []
            return "\n".join(text.strip() for text in texts if text and text.strip())
        except Exception:
            return ""

    def _parse_pdf(self, file_data: bytes, filename: str = "") -> Tuple[str, List]:
        """解析 PDF 文字层，并对空文字页做受限 OCR。"""
        import config
        import logging

        try:
            import pdfplumber

            logging.getLogger("pdfplumber").setLevel(logging.ERROR)

            text_parts = []
            all_tables = []
            ocr_document = None
            ocr_count = 0

            with pdfplumber.open(io.BytesIO(file_data)) as pdf:
                total_pages = len(pdf.pages)
                page_indices = self._select_pdf_page_indices(total_pages, filename)
                for i in page_indices:
                    page = pdf.pages[i]
                    # 仅提取文本，跳过表格提取（表格提取极慢且 HTML 中通常已有主要表格）
                    page_text = (page.extract_text() or "").strip()

                    if (
                        config.OCR_ENABLED
                        and len(page_text) < config.OCR_MIN_TEXT_CHARS
                        and ocr_count < config.OCR_MAX_PAGES_PER_PDF
                    ):
                        if ocr_document is None:
                            import pypdfium2 as pdfium
                            ocr_document = pdfium.PdfDocument(file_data)
                        ocr_text = self._ocr_pdf_page(ocr_document, i)
                        ocr_count += 1
                        if len(ocr_text) > len(page_text):
                            page_text = ocr_text

                    if page_text:
                        text_parts.append(f"[PDF 第 {i + 1} 页]\n{page_text}")

                if total_pages > len(page_indices):
                    shown = ",".join(str(i + 1) for i in page_indices)
                    text_parts.append(
                        f"[PDF 共 {total_pages} 页，已抽取第 {shown} 页]"
                    )

            if ocr_document is not None:
                ocr_document.close()

            return "\n".join(text_parts), all_tables

        except ImportError:
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(io.BytesIO(file_data))
                text_parts = []
                page_indices = self._select_pdf_page_indices(len(reader.pages), filename)
                for i in page_indices:
                    page = reader.pages[i]
                    text = page.extract_text()
                    if text:
                        text_parts.append(f"[PDF 第 {i + 1} 页]\n{text}")
                return "\n".join(text_parts), []
            except Exception:
                return "", []

    def _parse_pptx(self, file_data: bytes) -> Tuple[str, List]:
        """解析 PPT 文件"""
        from pptx import Presentation

        prs = Presentation(io.BytesIO(file_data))
        text_parts = []

        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text.strip())
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_data = [cell.text.strip() for cell in row.cells]
                        text_parts.append(" | ".join(row_data))

        return "\n".join(text_parts), []

    def _parse_txt(self, file_data: bytes) -> Tuple[str, List]:
        """解析纯文本文件"""
        for encoding in ["utf-8", "gbk", "gb2312", "latin-1"]:
            try:
                text = file_data.decode(encoding)
                return text.strip(), []
            except (UnicodeDecodeError, LookupError):
                continue
        return file_data.decode("utf-8", errors="ignore").strip(), []
