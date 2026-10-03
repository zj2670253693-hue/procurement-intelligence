"""快速测试脚本：验证框架各模块是否正常工作"""
import sys
import json
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

def test_imports():
    print("=== 1. 测试模块导入 ===")
    from config import REQUIRED_FIELDS, FIELD_LABELS
    print(f"核心字段 ({len(REQUIRED_FIELDS)} 个):")
    for f in REQUIRED_FIELDS:
        print(f"  - {FIELD_LABELS[f]} ({f})")
    print("  ✅ 配置模块导入成功")

    from models.schemas import ExtractedEntity, ExtractionResult, AnnouncementContent
    print("  ✅ 数据模型模块导入成功")

    from data_loader.html_parser import HTMLAnnouncementParser
    from data_loader.attachment_parser import AttachmentParser
    print("  ✅ 数据加载模块导入成功")

    from extractor.prompts import build_extraction_prompt, build_extra_info_prompt
    print("  ✅ Prompt 模板模块导入成功")

    from output.writer import ResultWriter
    print("  ✅ 输出模块导入成功")

    from evaluation.metrics import Evaluator
    print("  ✅ 评测模块导入成功")

    from database.connection import DatabaseManager, DatabaseConfig
    from database.dao import DataRepository
    from database.schema import ALL_TABLES
    print(f"  ✅ 数据库模块导入成功（预定义 {len(ALL_TABLES)} 张表）")


def test_low_memory_worker_configuration():
    print("\n=== 2. 测试 2核4GB 并发配置 ===")
    import importlib
    import os
    import config

    env_name = "PROCESS_MAX_WORKERS"
    original = os.environ.get(env_name)
    try:
        os.environ.pop(env_name, None)
        importlib.reload(config)
        assert config.PROCESS_MAX_WORKERS == 2

        os.environ[env_name] = "0"
        importlib.reload(config)
        assert config.PROCESS_MAX_WORKERS == 1

        os.environ[env_name] = "-3"
        importlib.reload(config)
        assert config.PROCESS_MAX_WORKERS == 1

        os.environ[env_name] = "4"
        importlib.reload(config)
        assert config.PROCESS_MAX_WORKERS == 4

        os.environ[env_name] = "abc"
        try:
            importlib.reload(config)
        except ValueError:
            pass
        else:
            raise AssertionError("非法并发数必须被拒绝")
    finally:
        if original is None:
            os.environ.pop(env_name, None)
        else:
            os.environ[env_name] = original
        importlib.reload(config)

    print("  ✅ 默认并发为2，非正数钳制为1，非法字符串被拒绝")


def test_deployment_assets():
    print("\n=== 3. 测试生产部署资产 ===")
    deploy_dir = Path(__file__).resolve().parent.parent / "deploy"
    required = {
        "procurement-api.service",
        "nginx-procurement-app.conf",
        "mysql-low-memory.cnf",
        "procurement.env.example",
        "bootstrap-ubuntu.sh",
        "finalize-services.sh",
    }
    missing = sorted(name for name in required if not (deploy_dir / name).is_file())
    assert not missing, f"缺少部署文件: {', '.join(missing)}"

    service = (deploy_dir / "procurement-api.service").read_text(encoding="utf-8")
    assert "User=zhoujin" in service
    assert "WorkingDirectory=/srv/procurement-app/current/backend" in service
    assert "--host 127.0.0.1" in service
    assert "--workers 1" in service

    nginx = (deploy_dir / "nginx-procurement-app.conf").read_text(encoding="utf-8")
    assert "root /srv/procurement-app/current/frontend-dist" in nginx
    assert "location /api/" in nginx
    assert "proxy_pass http://127.0.0.1:8000" in nginx

    mysql = (deploy_dir / "mysql-low-memory.cnf").read_text(encoding="utf-8")
    assert "bind-address = 127.0.0.1" in mysql
    assert "innodb_buffer_pool_size = 256M" in mysql
    assert "max_connections = 30" in mysql

    bootstrap = (deploy_dir / "bootstrap-ubuntu.sh").read_text(encoding="utf-8")
    assert "'procurement_app'@'127.0.0.1'" in bootstrap
    print("  ✅ systemd、Nginx、MySQL 和部署脚本约束正确")


def test_html_parser():
    print("\n=== 2. 测试 HTML 解析 ===")
    from data_loader.html_parser import HTMLAnnouncementParser
    parser = HTMLAnnouncementParser()
    content = parser.parse_file("sample_data/sample_announcement.html")
    print(f"  公告ID: {content.announcement_id}")
    print(f"  标题: {content.title[:40]}...")
    print(f"  文本长度: {len(content.html_text)} 字符")
    print(f"  表格数量: {len(content.html_tables)}")
    for i, table in enumerate(content.html_tables):
        print(f"  表格{i+1}: {len(table)} 行")
        if table:
            print(f"    表头: {table[0][:3]}...")
    print("  ✅ HTML 解析正常")


def test_evaluation():
    print("\n=== 3. 测试评测模块 ===")
    from evaluation.metrics import Evaluator
    from models.schemas import ExtractedEntity, ExtractionResult

    evaluator = Evaluator()

    # 模拟预测结果（部分正确）
    predictions = [
        ExtractionResult(
            announcement_id="sample_announcement",
            entities=[
                ExtractedEntity(
                    product_name="应用服务器", category="A02010103 服务器",
                    brand="华为", spec_model="2288H V5",
                    unit_price="35,000.00", quantity="10台", total_price="350,000.00"
                ),
                ExtractedEntity(
                    product_name="数据库服务器", category="A02010103 服务器",
                    brand="华为", spec_model="2288H V5",
                    unit_price="58,000.00", quantity="5台", total_price="290,000.00"
                ),
            ]
        )
    ]

    with open("sample_data/ground_truth_sample.json", "r", encoding="utf-8") as f:
        ground_truths = json.load(f)

    # 故意加入一篇完全缺失的预测，验证它会计入 FN，而不是被评测器忽略。
    ground_truths.append({
        "announcement_id": "missing_announcement",
        "entities": [{"product_name": "路由器"}],
    })

    metrics = evaluator.evaluate(predictions, ground_truths)
    assert metrics.total_samples == 2
    assert metrics.recall < 1.0
    assert not evaluator._field_match("服务器1", "路由器1", "product_name")
    assert evaluator._field_match("1.5万元", "15000元", "total_price")
    evaluator.print_report(metrics)
    print("  ✅ 评测模块正常")


def test_database_schema():
    print("\n=== 5. 测试数据库表结构 ===")
    from database.schema import (
        CREATE_TABLE_ANNOUNCEMENTS, CREATE_TABLE_ENTITIES,
        CREATE_TABLE_PROJECTS, CREATE_TABLE_BIDDERS, ALL_TABLES
    )
    print(f"  公告表 announcements: {'CREATE TABLE' in CREATE_TABLE_ANNOUNCEMENTS}")
    print(f"  标的物表 extracted_entities: {'CREATE TABLE' in CREATE_TABLE_ENTITIES}")
    print(f"  项目关系表 project_relations: {'CREATE TABLE' in CREATE_TABLE_PROJECTS}")
    print(f"  投标人表 bidders: {'CREATE TABLE' in CREATE_TABLE_BIDDERS}")
    print(f"  共 {len(ALL_TABLES)} 张表的 DDL 已定义")
    print("  ✅ 数据库表结构定义正常")


def test_prompts():
    print("\n=== 4. 测试 Prompt 模板 ===")
    from extractor.prompts import build_extraction_prompt, build_extra_info_prompt
    sys_prompt, user_prompt = build_extraction_prompt("测试文本", has_tables=False)
    print(f"  System Prompt 长度: {len(sys_prompt)} 字符")
    print(f"  User Prompt 包含测试文本: {'测试文本' in user_prompt}")
    _, table_prompt = build_extraction_prompt("正文和表格", has_tables=True)
    assert "正文和表格" in table_prompt
    assert "完整内容中含有" in table_prompt
    assert "以下是公告中提取到的表格内容" not in table_prompt
    print("  ✅ Prompt 模板正常")


def test_entity_deduplication():
    print("\n=== 6. 测试重复标的物去重 ===")
    from extractor.llm_extractor import LLMExtractor
    from models.schemas import ExtractedEntity

    entities = [
        ExtractedEntity(product_name="服务器", brand="华为", quantity="1 台"),
        ExtractedEntity(product_name="服务器", brand="华为", quantity="1台"),
        ExtractedEntity(product_name="服务器", brand="华为", quantity="2台"),
    ]
    deduplicated = LLMExtractor._deduplicate_entities(entities)
    assert len(deduplicated) == 2
    print("  ✅ 完全重复行已删除，数量不同的真实记录被保留")


def test_dataset_zip_detection():
    print("\n=== 7. 测试官方双压缩包识别 ===")
    from api.routers.task3 import _classify_zip, _extract_dataset_archive

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        html_archive = root / "html.zip"
        attachment_archive = root / "file.zip"
        single_attachment = root / "one.zip"
        output = root / "out"
        output.mkdir()

        with zipfile.ZipFile(html_archive, "w") as zf:
            zf.writestr("nested/a.html", "<html>公告</html>")
        with zipfile.ZipFile(attachment_archive, "w") as zf:
            zf.writestr("a.zip", b"a")
            zf.writestr("b.zip", b"b")
        with zipfile.ZipFile(single_attachment, "w") as zf:
            zf.writestr("报价明细.pdf", b"pdf")

        assert _classify_zip(html_archive) == "dataset"
        assert _classify_zip(attachment_archive) == "dataset"
        assert _classify_zip(single_attachment) == "attachment"
        assert _extract_dataset_archive(html_archive, output) == 1
        assert (output / "a.html").exists()
    print("  ✅ HTML 外层包、附件外层包与单篇附件包可正确区分")


def test_pipeline_attachment_tables_and_failure_callback():
    print("\n=== 8. 测试流水线附件表格与异常进度 ===")
    from openpyxl import Workbook
    import config
    from data_loader.html_parser import HTMLAnnouncementParser
    from models.schemas import ExtractionResult
    from pipeline import EntityExtractionPipeline

    class FakeExtractor:
        def __init__(self):
            self.has_tables = False

        def extract(self, text, has_tables=False, announcement_id=""):
            self.has_tables = has_tables
            return ExtractionResult(announcement_id=announcement_id)

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        html = root / "notice.html"
        html.write_text("<html><body><p>公告正文</p></body></html>", encoding="utf-8")

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["采购标的", "数量"])
        sheet.append(["服务器", 1])
        excel_data = BytesIO()
        workbook.save(excel_data)
        with zipfile.ZipFile(root / "notice.zip", "w") as zf:
            zf.writestr("报价明细.xlsx", excel_data.getvalue())

        pipeline = EntityExtractionPipeline.__new__(EntityExtractionPipeline)
        pipeline.html_parser = HTMLAnnouncementParser()
        pipeline.extractor = FakeExtractor()
        pipeline.extract_extra = False
        pipeline.repository = None
        result = pipeline.process_single_file(str(html))
        assert result.success
        assert pipeline.extractor.has_tables

        # 模拟 worker 抛异常，进度回调必须收到本次失败结果，不能引用上一条结果。
        pipeline.max_workers = 1
        pipeline.process_single_file = lambda _path: (_ for _ in ()).throw(RuntimeError("boom"))
        callbacks = []
        original_output_dir = config.OUTPUT_DIR
        config.OUTPUT_DIR = root
        try:
            results = pipeline.process_directory(
                str(root),
                progress_callback=lambda done, total, success, item: callbacks.append(item),
            )
        finally:
            config.OUTPUT_DIR = original_output_dir
        assert len(results) == 1 and not results[0].success
        assert callbacks[0].announcement_id == "notice"
    print("  ✅ 附件表格进入 Prompt 状态，worker 异常可正确上报")


def test_pipeline_checkpoint_resume():
    print("\n=== 9. 测试全量处理检查点与断点续跑 ===")
    import json
    import config
    from models.schemas import ExtractionResult, ExtractedEntity
    from output.writer import ResultWriter
    from pipeline import EntityExtractionPipeline

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        input_dir = root / "input"
        output_dir = root / "output"
        input_dir.mkdir()
        for index in range(10):
            (input_dir / f"notice_{index:02d}.html").write_text(
                "<html><body>公告</body></html>", encoding="utf-8"
            )

        pipeline = EntityExtractionPipeline.__new__(EntityExtractionPipeline)
        pipeline.writer = ResultWriter(output_dir)
        pipeline.max_workers = 1
        pipeline.process_single_file = lambda path: ExtractionResult(
            announcement_id=Path(path).stem,
            entities=[ExtractedEntity(product_name="测试标的")],
            success=True,
        )
        previous = [ExtractionResult(announcement_id="previous", success=True)]

        original_output_dir = config.OUTPUT_DIR
        config.OUTPUT_DIR = output_dir
        try:
            results = pipeline.process_directory(
                str(input_dir), checkpoint_results=previous
            )
            checkpoint_path = output_dir / "raw_results.checkpoint.json"
            assert checkpoint_path.exists()
            checkpoint_data = json.loads(checkpoint_path.read_text(encoding="utf-8"))
            assert len(results) == 10
            assert len(checkpoint_data) == 11

            loaded, skip_ids = pipeline._load_previous_results()
            assert len(loaded) == 11
            assert "previous" in skip_ids and "notice_09" in skip_ids

            pipeline._clear_checkpoint()
            assert not checkpoint_path.exists()
            pipeline.writer.write_raw_results(previous)
            loaded, skip_ids = pipeline._load_previous_results()
            assert len(loaded) == 1 and skip_ids == {"previous"}
        finally:
            config.OUTPUT_DIR = original_output_dir
    print("  ✅ 每 10 条自动保存，续跑优先读取检查点，完成后可清理")


def test_attachment_content_detection_and_pdf_sampling():
    print("\n=== 10. 测试 .doc 真实格式识别与长 PDF 抽页 ===")
    from docx import Document
    from data_loader.attachment_parser import AttachmentParser

    parser = AttachmentParser()

    # 后缀名是 .doc，实际容器是 DOCX。
    document = Document()
    document.add_paragraph("伪装后缀的 DOCX 正文")
    docx_data = BytesIO()
    document.save(docx_data)
    text, _ = parser._parse_file(docx_data.getvalue(), ".doc", "wrong.doc")
    assert "伪装后缀的 DOCX 正文" in text

    # 后缀名是 .doc，实际内容是 HTML。
    text, _ = parser._parse_file(
        "<html><body><table><tr><td>报价明细</td></tr></table></body></html>".encode("utf-8"),
        ".doc",
        "html.doc",
    )
    assert "报价明细" in text

    salvage_source = "这是一段用于验证旧版 Word 损坏文件文本抢救能力的中文正文内容。"
    salvaged = parser._salvage_doc_unicode(salvage_source.encode("utf-16le"))
    assert "文本抢救能力" in salvaged

    # 普通长 PDF 必须覆盖首页、中间页和末页，不再只读前 15 页。
    indices = parser._select_pdf_page_indices(100, "采购文件.pdf")
    assert indices[0] == 0 and indices[-1] == 99
    assert any(30 <= i <= 70 for i in indices)
    assert len(indices) <= 20

    # 报价类关键附件允许覆盖更多页。
    relevant_indices = parser._select_pdf_page_indices(100, "分项报价明细.pdf")
    assert relevant_indices[-1] == 99
    assert len(relevant_indices) > len(indices)
    print("  ✅ 伪 DOCX/HTML 可按真实格式解析，长 PDF 覆盖首中尾")


if __name__ == "__main__":
    try:
        test_imports()
        test_low_memory_worker_configuration()
        test_deployment_assets()
        test_html_parser()
        test_evaluation()
        test_prompts()
        test_database_schema()
        test_entity_deduplication()
        test_dataset_zip_detection()
        test_pipeline_attachment_tables_and_failure_callback()
        test_pipeline_checkpoint_resume()
        test_attachment_content_detection_and_pdf_sampling()
        print("\n" + "=" * 50)
        print("🎉 所有测试通过！框架基础功能正常。")
        print("=" * 50)
        print("\n下一步：")
        print("  1. 配置 .env 中的 API Key 和数据库信息")
        print("  2. 运行: python main.py --input sample_data/sample_announcement.html")
    except Exception as e:
        print(f"\n❌ 测试失败: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
